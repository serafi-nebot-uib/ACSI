#!/usr/bin/env python3

import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook

wb = load_workbook("data.xlsx")
ws = wb.active  # or wb["SheetName"]

rows = ws.iter_rows(values_only=True)
header = next(rows)

req_sec, proc_time, queue_time, send_time = (np.array([b for b in a if b], dtype="float32") for a in zip(*rows))
assert req_sec.size == proc_time.size == queue_time.size == send_time.size, "data sets length mismatch"
total_time = np.array([sum(a) for a in zip(proc_time, queue_time, send_time)])
N = len(total_time)

def line():
  x = np.arange(0, req_sec.size, 1)

  fig, ax_time = plt.subplots()

  ax_time.set_ylabel("time (s)")
  total_time_line, *_ = ax_time.plot(x, total_time, label="TotalTime", color="black")
  proc_time_line, *_ = ax_time.plot(x, proc_time, label="ProessTime", color="red")
  queue_time_line, *_ = ax_time.plot(x, queue_time, label="QueueTime", color="blue")
  send_time_line, *_ = ax_time.plot(x, send_time, label="SendTime", color="green")

  ax_prod = ax_time.twinx()
  ax_prod.set_ylabel("productivity (req/sec)")
  req_sec_line, *_ = ax_prod.plot(x, req_sec, label="requests/s", color="orange")

  lines = [total_time_line, proc_time_line, queue_time_line, send_time_line, req_sec_line]
  labels = [l.get_label() for l in lines]
  ax_time.legend(lines, labels, loc="upper right")

  fig.tight_layout()
  plt.show()

def scatter():
  fig, ax = plt.subplots()

  ax.set_xlabel("productivity (req/sec)")
  ax.set_ylabel("time (s)")
  # ax.scatter(req_sec, total_time, label="TotalTime", color="black")
  ax.scatter(req_sec, proc_time, label="ProessTime", color="red")
  # ax.scatter(req_sec, queue_time, label="QueueTime", color="blue")
  # ax.scatter(req_sec, send_time, label="SendTime", color="green")

  # datasets = [proc_time_line, queue_time_line, send_time_line]
  # labels = [l.get_label() for l in datasets]
  ax.legend(loc="upper right")

  fig.tight_layout()
  plt.show()

# line()
# scatter()

x = req_sec
y = proc_time

dist = np.zeros((N, N))

for i in range(N):
  for j in range(i+1):
    dist[j][i] = np.sqrt((x[i]-x[j])**2 + (y[i]-y[j])**2)

print(dist)
