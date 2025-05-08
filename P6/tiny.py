#!/usr/bin/env python3

from tinygrad import Tensor, dtypes
from openpyxl import load_workbook

wb = load_workbook("data.xlsx")
ws = wb.active  # or wb["SheetName"]

rows = ws.iter_rows(values_only=True)
header = next(rows)

req_sec, proc_time, queue_time, send_time = ([b for b in a if b] for a in zip(*rows))
assert len(req_sec) == len(proc_time) == len(queue_time) == len(send_time), "data sets length mismatch"
total_time = [sum(a) for a in zip(proc_time, queue_time, send_time)]
N = len(total_time)

x = Tensor(req_sec, dtype=dtypes.float32)
y = Tensor(total_time, dtype=dtypes.float32)

diff_x = (x.unsqueeze(0) - x.unsqueeze(1)).triu()
diff_y = (y.unsqueeze(0) - y.unsqueeze(1)).triu()
dist = (diff_x**2 + diff_y**2).sqrt()

m = dist.numpy()
mval = m[m != 0].min()

mcoords = ((i, j) for i in range(N-1) for j in range(i+1, N) if abs(m[i][j] - mval) < 0.001)
for i, j in mcoords: print(f"m[{i}][{j}]: {m[i][j]}")

# print(r.numpy())
# print(r.min().numpy())
# print(r.max().numpy())
