#!/usr/bin/env python3

import numpy as np
import matplotlib.pyplot as plt
from itertools import accumulate
from openpyxl import load_workbook

dtype = "float32"

def linear_regression(x, y):
  assert x.shape == y.shape, "x and y axis must have the sape shape"
  assert len(x.shape) == 1 == len(x.shape), "input data must only have 1 dimension"
  n, xm, ym = x.shape[0], np.mean(x), np.mean(y)
  b = (np.sum(x * y) - n * xm * ym) / (np.sum(x**2) - n * xm**2)
  a = ym - b * xm
  return a, b

def ma(y, n):
  assert len(y.shape) == 1, "input data must only have 1 dimension"
  return np.fromiter((np.mean(y[i-n:i]) for i in range(n, y.shape[0] + 1)), like=y, dtype=dtype)

def exp(y, alpha):
  assert len(y.shape) == 1, "input data must only have 1 dimension"
  assert 0 < alpha < 1, "alpha must be 0 < alpha < 1"
  return np.fromiter(accumulate(y[1:], lambda a, b: a+alpha*(b-a), initial=y[0]), like=y, dtype=dtype)

wb = load_workbook("data/data.xlsx")
ws = wb.active  # or wb["SheetName"]

rows = ws.iter_rows(values_only=True)
header = next(rows)

req_sec, proc_time, queue_time, send_time = (np.array([b for b in a if b], dtype=dtype) for a in zip(*rows))
assert req_sec.size == proc_time.size == queue_time.size == send_time.size, "data sets length mismatch"
total_time = np.array([sum(a) for a in zip(proc_time, queue_time, send_time)])
N = len(total_time)

def line():
  x = np.arange(0, req_sec.size, 1)

  fig, ax_time = plt.subplots(figsize=(12, 8))

  ax_time.set_ylabel("time (s)")
  # total_time_line, *_ = ax_time.plot(x, total_time, label="TotalTime", color="black")
  # queue_time_line, *_ = ax_time.plot(x, queue_time, label="QueueTime", color="blue")
  # proc_time_line, *_ = ax_time.plot(x, proc_time, label="ProessTime", color="red")
  send_time_line, *_ = ax_time.plot(x, send_time, label="SendTime", color="green")

  # ax_prod = ax_time.twinx()
  # ax_prod.set_xlabel("measure number")
  # ax_prod.set_ylabel("productivity (req/sec)")

  # req_sec_line, *_ = ax_prod.plot(x, req_sec, label="requests/s", color="orange")

  # lines = [proc_time_line, queue_time_line, send_time_line, req_sec_line]
  lines = [send_time_line]
  labels = [l.get_label() for l in lines]
  ax_time.legend(lines, labels, loc="upper right")

  fig.tight_layout()
  plt.show()

def error(real, pred):
  assert real.shape == pred.shape, "x and y axis must have the sape shape"
  assert len(real.shape) == 1 == len(pred.shape), "input data must only have 1 dimension"
  return np.sum((real - pred)**2) / real.shape[0]

def model(y, name="y"):
  x = np.arange(0, N, 1)

  a, b = linear_regression(x, y)
  lingres = a + b * x
  lingres_err = error(y, lingres)
  lingres_pred = a + b * N

  ma70, ma500, ma2000 = ma(y, 70), ma(y, 500), ma(y, 2000)
  ma70_err = error(y[70:], ma70[:len(ma70)-1])
  ma500_err = error(y[500:], ma500[:len(ma500)-1])
  ma2000_err = error(y[2000:], ma2000[:len(ma2000)-1])

  exp_y = exp(y, alpha=0.6)
  exp_err = error(y, exp_y)

  print(name)
  print(f"linear regression: prediction = {lingres_pred:.12f} error = {lingres_err:.12f}")
  print(f"             ma70: prediction = {ma70[-1]:.12f} error = {ma70_err:.12f}")
  print(f"            ma500: prediction = {ma500[-1]:.12f} error = {ma500_err:.12f}")
  print(f"           ma2000: prediction = {ma2000[-1]:.12f} error = {ma2000_err:.12f}")
  print(f"   exp. smoothing: prediction = {exp_y[-1]:.12f} error = {exp_err:.12f}")
  print()
  return

  fig, ax = plt.subplots(figsize=(12, 8))

  ax.set_xlabel("measure number")
  ax.set_ylabel(name)
  ax.scatter(x, y, label="real", color="lightgray")
  # ax.plot(x, y, label="y")

  ax.plot(x, exp_y, color="blue", label="exp. smooth. (α=0.6)")
  ax.plot(x[70:], ma70[:len(ma70)-1], color="orange", label="MA 70", linewidth=2)
  ax.plot(x[500:], ma500[:len(ma500)-1], color="green", label="MA 500", linewidth=2)
  ax.plot(x[2000:], ma2000[:len(ma2000)-1], color="magenta", label="MA 2000", linewidth=2)
  ax.plot(x, lingres, color="red", label="linear regression", linewidth=2)

  fig.tight_layout()
  plt.legend()
  plt.show()

line()
# model(req_sec, "requests/s")
# model(proc_time, "ProcessTime")
# model(queue_time, "QueueTime")
# model(send_time, "SendTime")
